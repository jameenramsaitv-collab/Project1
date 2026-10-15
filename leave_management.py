"""Excel-backed leave-management MCP tools for the Project1 HR support server.

Prototype note:
- Employee names and manager emails are demo identifiers, not authentication.
- Before production, bind the employee to a verified login/session and verify manager
  approval through authenticated links or an authenticated manager workflow.
"""
from __future__ import annotations

import os
import re
import smtplib
import uuid
from datetime import date, datetime, timedelta
from email.message import EmailMessage
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

BASE_DIR = Path(__file__).resolve().parent
WORKBOOK_PATH = Path(os.getenv("HR_WORKBOOK_PATH", str(BASE_DIR / "hr_data.xlsx")))

EMPLOYEE_HEADERS = ["employee_id", "name", "email", "manager_id"]
MANAGER_HEADERS = ["manager_id", "name", "email"]
BALANCE_HEADERS = ["employee_id", "leave_type", "entitled", "used"]
REQUEST_HEADERS = [
    "request_id", "employee_id", "leave_type", "start_date", "end_date",
    "days", "reason", "status", "created_at", "manager_email", "decision_at",
    "decision_note",
]


def ensure_workbook() -> None:
    """Create the workbook and required sheets if they do not already exist."""
    WORKBOOK_PATH.parent.mkdir(parents=True, exist_ok=True)
    if WORKBOOK_PATH.exists():
        wb = load_workbook(WORKBOOK_PATH)
    else:
        wb = Workbook()
        wb.active.title = "Employees"
        wb.create_sheet("Managers")
        wb.create_sheet("LeaveBalances")
        wb.create_sheet("LeaveRequests")

    expected = {
        "Employees": EMPLOYEE_HEADERS,
        "Managers": MANAGER_HEADERS,
        "LeaveBalances": BALANCE_HEADERS,
        "LeaveRequests": REQUEST_HEADERS,
    }
    for sheet_name, headers in expected.items():
        if sheet_name not in wb.sheetnames:
            wb.create_sheet(sheet_name)
        ws = wb[sheet_name]
        if ws.max_row == 1 and all(ws.cell(1, i + 1).value is None for i in range(len(headers))):
            for col, header in enumerate(headers, start=1):
                ws.cell(1, col, header)
        elif ws.max_row == 1 and ws.cell(1, 1).value is None:
            for col, header in enumerate(headers, start=1):
                ws.cell(1, col, header)
    wb.save(WORKBOOK_PATH)
    wb.close()


def _rows(sheet_name: str) -> list[dict[str, Any]]:
    ensure_workbook()
    wb = load_workbook(WORKBOOK_PATH, data_only=True)
    try:
        ws = wb[sheet_name]
        values = list(ws.iter_rows(values_only=True))
        if not values:
            return []
        headers = [str(h).strip() if h is not None else "" for h in values[0]]
        result = []
        for row in values[1:]:
            if not any(value is not None for value in row):
                continue
            result.append({headers[i]: row[i] if i < len(row) else None for i in range(len(headers)) if headers[i]})
        return result
    finally:
        wb.close()


def _append(sheet_name: str, values: list[Any]) -> None:
    ensure_workbook()
    wb = load_workbook(WORKBOOK_PATH)
    try:
        wb[sheet_name].append(values)
        wb.save(WORKBOOK_PATH)
    finally:
        wb.close()


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _find_employee(employee_name: str) -> dict[str, Any] | None:
    target = employee_name.strip().casefold()
    matches = [r for r in _rows("Employees") if _text(r.get("name")).casefold() == target]
    if len(matches) > 1:
        raise ValueError("More than one employee has that name. Use a unique employee name or employee ID.")
    return matches[0] if matches else None


def _find_employee_by_id(employee_id: str) -> dict[str, Any] | None:
    target = employee_id.strip().casefold()
    matches = [r for r in _rows("Employees") if _text(r.get("employee_id")).casefold() == target]
    return matches[0] if matches else None


def _manager_for(employee: dict[str, Any]) -> dict[str, Any] | None:
    manager_id = _text(employee.get("manager_id"))
    return next((m for m in _rows("Managers") if _text(m.get("manager_id")).casefold() == manager_id.casefold()), None)


def _parse_date(value: str) -> date:
    try:
        return date.fromisoformat(value.strip())
    except (ValueError, AttributeError):
        raise ValueError(f"Invalid date '{value}'. Use YYYY-MM-DD.") from None


def _business_days(start: date, end: date) -> int:
    """Count Monday-Friday days inclusively; holidays are not yet modeled."""
    days = 0
    current = start
    while current <= end:
        if current.weekday() < 5:
            days += 1
        current += timedelta(days=1)
    return days


def _request_dates(row: dict[str, Any]) -> tuple[date, date]:
    start = row.get("start_date")
    end = row.get("end_date")
    if isinstance(start, datetime):
        start = start.date()
    elif isinstance(start, str):
        start = _parse_date(start[:10])
    if isinstance(end, datetime):
        end = end.date()
    elif isinstance(end, str):
        end = _parse_date(end[:10])
    if not isinstance(start, date) or not isinstance(end, date):
        raise ValueError("A stored leave request has an invalid date.")
    return start, end


def _leave_summary(employee: dict[str, Any]) -> list[dict[str, Any]]:
    employee_id = _text(employee.get("employee_id"))
    balances = [r for r in _rows("LeaveBalances") if _text(r.get("employee_id")).casefold() == employee_id.casefold()]
    requests = [r for r in _rows("LeaveRequests") if _text(r.get("employee_id")).casefold() == employee_id.casefold()]
    summary = []
    for balance in balances:
        leave_type = _text(balance.get("leave_type"))
        entitled = float(balance.get("entitled") or 0)
        used = float(balance.get("used") or 0)
        pending = 0
        for req in requests:
            if _text(req.get("leave_type")).casefold() != leave_type.casefold():
                continue
            if _text(req.get("status")).casefold() == "pending":
                pending += float(req.get("days") or 0)
        available = entitled - used - pending
        summary.append({
            "leave_type": leave_type,
            "entitled": entitled,
            "used": used,
            "pending": pending,
            "available": available,
        })
    return summary


def _archive_and_send(
    to_email: str, cc_email: str, subject: str, body: str
) -> str:
    """Archive the message and send it if SMTP is configured."""
    if not re.match(r"^[^@\s]+@[^@\s]+\.[^@\s]+$", to_email):
        raise ValueError("Manager email address is missing or invalid.")
    if not re.match(r"^[^@\s]+@[^@\s]+\.[^@\s]+$", cc_email):
        raise ValueError("Employee email address is missing or invalid.")

    archive_dir = BASE_DIR / "sent_emails"
    archive_dir.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    safe_subject = re.sub(r"[^A-Za-z0-9_-]+", "_", subject)[:50]
    archive_path = archive_dir / f"leave_{timestamp}_{safe_subject}.txt"
    archive_path.write_text(
        f"TO: {to_email}\nCC: {cc_email}\nSUBJECT: {subject}\n"
        f"TIMESTAMP: {datetime.now().isoformat(timespec='seconds')}\n\n{body}\n",
        encoding="utf-8",
    )

    host = os.getenv("SMTP_HOST", "").strip()
    port = int(os.getenv("SMTP_PORT", "587"))
    user = os.getenv("SMTP_USER", "").strip()
    password = os.getenv("SMTP_PASSWORD", "").strip()
    from_email = (os.getenv("SMTP_FROM_EMAIL") or user).strip()

    if not (host and user and password and from_email):
        return f"Approval email archived at {archive_path.name}; SMTP is not configured, so no email was sent."

    message = EmailMessage()
    message["Subject"] = subject
    message["From"] = from_email
    message["To"] = to_email
    message["Cc"] = cc_email
    message.set_content(body)

    if port == 465:
        with smtplib.SMTP_SSL(host, port, timeout=15) as server:
            server.login(user, password)
            server.send_message(message)
    else:
        with smtplib.SMTP(host, port, timeout=15) as server:
            server.starttls()
            server.login(user, password)
            server.send_message(message)
    return f"Approval email sent to manager and CC'd to employee. Archived as {archive_path.name}."


def register_leave_tools(mcp: Any) -> None:
    """Register the leave-management tools on the existing FastMCP instance."""
    ensure_workbook()

    @mcp.tool()
    def get_employee_leave_balance(employee_name: str) -> str:
        """Show entitlement, used, pending, and available leave for an employee by exact name."""
        try:
            employee = _find_employee(employee_name)
            if not employee:
                return f"No employee named '{employee_name}' was found in hr_data.xlsx."
            summary = _leave_summary(employee)
            if not summary:
                return f"No leave balances are configured for {_text(employee.get('name'))}."
            lines = [
                f"Leave summary for {_text(employee.get('name'))} ({_text(employee.get('employee_id'))}):"
            ]
            for item in summary:
                lines.append(
                    f"- {item['leave_type']}: entitled {item['entitled']:g}, "
                    f"used {item['used']:g}, pending {item['pending']:g}, "
                    f"available {item['available']:g} days"
                )
            return "\n".join(lines)
        except Exception as exc:
            return f"Could not retrieve leave balance: {exc}"

    @mcp.tool()
    def get_employee_leave_history(employee_name: str) -> str:
        """List leave requests for an employee, including dates, type, days, and status."""
        try:
            employee = _find_employee(employee_name)
            if not employee:
                return f"No employee named '{employee_name}' was found."
            employee_id = _text(employee.get("employee_id"))
            requests = [
                r for r in _rows("LeaveRequests")
                if _text(r.get("employee_id")).casefold() == employee_id.casefold()
            ]
            if not requests:
                return f"No leave requests found for {_text(employee.get('name'))}."
            lines = [f"Leave history for {_text(employee.get('name'))}:"]
            for req in requests:
                lines.append(
                    f"- {req.get('request_id')}: {req.get('leave_type')}, "
                    f"{req.get('start_date')} to {req.get('end_date')} "
                    f"({req.get('days')} working days), status: {req.get('status')}"
                )
            return "\n".join(lines)
        except Exception as exc:
            return f"Could not retrieve leave history: {exc}"

    @mcp.tool()
    def apply_for_leave(
        employee_name: str,
        leave_type: str,
        start_date: str,
        end_date: str,
        reason: str,
    ) -> str:
        """Create a pending leave request and email the employee's manager with the employee CC'd.

        Dates must be YYYY-MM-DD. Weekends are excluded from requested working days.
        The request is not approved automatically.
        """
        try:
            employee = _find_employee(employee_name)
            if not employee:
                return f"No employee named '{employee_name}' was found."
            employee_email = _text(employee.get("email"))
            if not employee_email:
                return "Employee email is missing in the Employees sheet."
            manager = _manager_for(employee)
            if not manager or not _text(manager.get("email")):
                return "No manager/email is configured for this employee. Request was not created."

            start = _parse_date(start_date)
            end = _parse_date(end_date)
            if start < date.today():
                return "Leave start date cannot be in the past."
            if end < start:
                return "End date must be on or after the start date."
            days = _business_days(start, end)
            if days <= 0:
                return "The selected dates contain no weekdays. Please choose valid leave dates."
            if not reason.strip():
                return "Please provide a reason for the leave request."

            balances = [
                b for b in _rows("LeaveBalances")
                if _text(b.get("employee_id")).casefold() == _text(employee.get("employee_id")).casefold()
                and _text(b.get("leave_type")).casefold() == leave_type.strip().casefold()
            ]
            if not balances:
                return f"Leave type '{leave_type}' is not configured for this employee."
            summary = next((s for s in _leave_summary(employee) if s["leave_type"].casefold() == leave_type.strip().casefold()), None)
            if not summary:
                return f"No balance exists for leave type '{leave_type}'."
            if days > summary["available"]:
                return (
                    f"Insufficient {leave_type} leave. Requested {days} working days, "
                    f"but only {summary['available']:g} days are available "
                    "(pending requests are reserved against the balance)."
                )

            # Avoid overlapping active requests for the same employee.
            for existing in _rows("LeaveRequests"):
                if _text(existing.get("employee_id")).casefold() != _text(employee.get("employee_id")).casefold():
                    continue
                if _text(existing.get("status")).casefold() not in {"pending", "approved"}:
                    continue
                old_start, old_end = _request_dates(existing)
                if start <= old_end and end >= old_start:
                    return f"These dates overlap an existing {existing.get('status')} request ({existing.get('request_id')})."

            request_id = "LR-" + uuid.uuid4().hex[:8].upper()
            manager_email = _text(manager.get("email"))
            subject = f"Leave approval requested: {request_id} — {_text(employee.get('name'))}"
            body = (
                "A leave request requires your decision.\n\n"
                f"Request ID: {request_id}\n"
                f"Employee: {_text(employee.get('name'))} ({_text(employee.get('employee_id'))})\n"
                f"Leave type: {leave_type.strip()}\n"
                f"Dates: {start.isoformat()} to {end.isoformat()}\n"
                f"Working days requested: {days}\n"
                f"Reason: {reason.strip()}\n\n"
                "Current status: PENDING\n"
                "Please approve or reject this request through the authorized HR workflow. "
                "Do not treat this notification as an approval."
            )
            _append("LeaveRequests", [
                request_id, _text(employee.get("employee_id")), leave_type.strip(),
                start.isoformat(), end.isoformat(), days, reason.strip(), "Pending",
                datetime.now().isoformat(timespec="seconds"), manager_email, "", "",
            ])
            try:
                notification = _archive_and_send(manager_email, employee_email, subject, body)
            except Exception as exc:
                notification = f"Request saved as Pending, but email delivery failed: {exc}"
            return (
                f"Leave request {request_id} created with status PENDING for {days} working day(s). "
                f"{notification} The request is not approved yet."
            )
        except Exception as exc:
            return f"Could not apply for leave: {exc}"

    @mcp.tool()
    def get_leave_request_status(request_id: str) -> str:
        """Find a leave request by its ID and return its current status."""
        request = next(
            (r for r in _rows("LeaveRequests") if _text(r.get("request_id")).casefold() == request_id.strip().casefold()),
            None,
        )
        if not request:
            return f"Leave request '{request_id}' was not found."
        return (
            f"Request {request.get('request_id')}: {request.get('status')}; "
            f"employee ID {request.get('employee_id')}; {request.get('leave_type')} "
            f"from {request.get('start_date')} to {request.get('end_date')}; "
            f"decision note: {_text(request.get('decision_note')) or 'None'}"
        )

    @mcp.tool()
    def record_leave_decision(
        request_id: str,
        manager_email: str,
        decision: str,
        decision_note: str = "",
    ) -> str:
        """Record a manager decision for a leave request.

        DEMO ONLY: manager_email is an asserted value, not authenticated identity.
        Do not expose this tool to untrusted users in production.
        """
        normalized_decision = decision.strip().casefold()
        if normalized_decision not in {"approve", "approved", "reject", "rejected"}:
            return "Decision must be approve or reject."
        approved = normalized_decision in {"approve", "approved"}
        ensure_workbook()
        wb = load_workbook(WORKBOOK_PATH)
        try:
            ws = wb["LeaveRequests"]
            headers = [cell.value for cell in ws[1]]
            index = {str(name): i + 1 for i, name in enumerate(headers) if name is not None}
            row_num = None
            row_values = None
            for r in range(2, ws.max_row + 1):
                if _text(ws.cell(r, index["request_id"]).value).casefold() == request_id.strip().casefold():
                    row_num = r
                    row_values = {key: ws.cell(r, col).value for key, col in index.items()}
                    break
            if row_num is None or row_values is None:
                return f"Leave request '{request_id}' was not found."
            if _text(row_values.get("status")).casefold() != "pending":
                return f"Request {request_id} is already {row_values.get('status')} and cannot be changed."
            employee = _find_employee_by_id(_text(row_values.get("employee_id")))
            manager = _manager_for(employee) if employee else None
            expected_email = _text(manager.get("email")) if manager else ""
            if not expected_email or expected_email.casefold() != manager_email.strip().casefold():
                return "The supplied manager email does not match the employee's assigned manager."
            new_status = "Approved" if approved else "Rejected"
            ws.cell(row_num, index["status"], new_status)
            ws.cell(row_num, index["decision_at"], datetime.now().isoformat(timespec="seconds"))
            ws.cell(row_num, index["decision_note"], decision_note.strip())
            # Update used leave only when approved; avoid double-counting by relying on status.
            if approved:
                employee_id = _text(row_values.get("employee_id"))
                leave_type = _text(row_values.get("leave_type"))
                days = float(row_values.get("days") or 0)
                balances_ws = wb["LeaveBalances"]
                b_headers = [cell.value for cell in balances_ws[1]]
                bidx = {str(name): i + 1 for i, name in enumerate(b_headers) if name is not None}
                for br in range(2, balances_ws.max_row + 1):
                    if (
                        _text(balances_ws.cell(br, bidx["employee_id"]).value).casefold() == employee_id.casefold()
                        and _text(balances_ws.cell(br, bidx["leave_type"]).value).casefold() == leave_type.casefold()
                    ):
                        current_used = float(balances_ws.cell(br, bidx["used"]).value or 0)
                        balances_ws.cell(br, bidx["used"], current_used + days)
                        break
            wb.save(WORKBOOK_PATH)
            employee_email = _text(employee.get("email")) if employee else ""
        finally:
            wb.close()

        # Best-effort notification; decision remains recorded even if SMTP fails.
        if employee_email:
            try:
                _archive_and_send(
                    employee_email,
                    expected_email,
                    f"Leave request {request_id} {new_status.lower()}",
                    f"Your leave request {request_id} has been {new_status.lower()}.\n"
                    f"Decision note: {decision_note.strip() or 'No note provided.'}\n",
                )
            except Exception:
                pass
        return f"Request {request_id} marked {new_status}. Leave balance updated if approved."
